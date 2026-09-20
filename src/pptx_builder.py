"""
Builds the daily status deck from report_data using python-pptx.
Keeps the "Situation Room" palette from the original design.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

NAVY = RGBColor(0x1E, 0x27, 0x61)
NAVY_DK = RGBColor(0x14, 0x1B, 0x47)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x5A, 0x6B, 0x8C)
PAPER = RGBColor(0xF4, 0xF6, 0xFB)
INK = RGBColor(0x1B, 0x23, 0x40)
MUTED = RGBColor(0x6B, 0x75, 0x91)
AMBER = RGBColor(0xD9, 0x8A, 0x2B)
RED = RGBColor(0xB3, 0x26, 0x1E)
GREEN = RGBColor(0x2E, 0x7D, 0x5B)

FONT_HEAD = "Cambria"
FONT_BODY = "Calibri"


def _bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _box(slide, x, y, w, h, text, size, color, bold=False, italic=False,
         font=FONT_BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = color
    return tb


def _card(slide, x, y, w, h, fill_color, line_color=None):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))  # 1 = rectangle
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def build_deck(report_data: dict, output_path: str):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    _slide_dashboard(prs.slides.add_slide(blank), report_data)
    _slide_event_log(prs.slides.add_slide(blank), report_data)
    _slide_trend(prs.slides.add_slide(blank), report_data)
    _slide_narrative_stability(prs.slides.add_slide(blank), report_data)
    _slide_source_library(prs.slides.add_slide(blank), report_data)
    _slide_deception_watch(prs.slides.add_slide(blank), report_data)
    _slide_proposal(prs.slides.add_slide(blank), report_data)

    prs.save(output_path)


def _slide_dashboard(s, d):
    _bg(s, NAVY_DK)
    _box(s, 0.6, 0.4, 12.1, 0.4, "RUSSIA \u2192 NATO / EU / UKRAINE", 13, ICE, bold=True, font=FONT_BODY)
    _box(s, 0.6, 0.8, 12.1, 0.7, "Hybrid Threat & Escalation Monitor", 30, WHITE, bold=True, font=FONT_HEAD)
    _box(s, 0.6, 1.45, 12.1, 0.35, f"Report date: {d['date']}  |  24h monitoring, consolidated 05:00 CET", 12, SLATE, italic=True)

    _card(s, 0.6, 2.0, 3.9, 3.1, NAVY)
    _box(s, 0.9, 2.2, 3.3, 0.35, "NATO / EU\u2013RUSSIA THREAT INDEX", 11, ICE, bold=True)
    _box(s, 0.9, 2.5, 2.2, 1.2, f"{d['threat_index']:.1f}", 70, AMBER, bold=True, font=FONT_HEAD)
    _box(s, 0.9, 3.75, 3.3, 0.3, f"Trend: {d['trend_label']}", 12, AMBER, bold=True)
    _box(s, 0.9, 4.15, 3.3, 0.85, d["threat_reading"], 10.5, ICE)

    _card(s, 4.75, 2.0, 7.95, 1.35, NAVY)
    _box(s, 5.0, 2.2, 7.5, 0.3, "SELF-CHECK: YESTERDAY'S ASSESSMENTS", 10.5, AMBER, bold=True)
    _box(s, 5.0, 2.55, 7.5, 0.7, d["self_check_summary"], 12, WHITE)

    _card(s, 4.75, 3.5, 7.95, 1.6, NAVY)
    _box(s, 5.0, 3.7, 7.5, 0.3, "TOP DEVELOPMENT, LAST 24H", 10.5, AMBER, bold=True)
    _box(s, 5.0, 4.05, 7.5, 1.0, d["top_development"], 12, WHITE)

    _box(s, 0.6, 7.1, 12.1, 0.3, "AUTOMATED OSINT SYNTHESIS - NOT FOR OPERATIONAL DECISION-MAKING", 9, SLATE)


def _slide_event_log(s, d):
    _bg(s, PAPER)
    _box(s, 0.6, 0.4, 8, 0.35, "LAST-24H EVENT LOG", 12, AMBER, bold=True)
    _box(s, 0.6, 0.72, 12.1, 0.55, "Cross-Referenced Developments & Credibility Scoring", 24, INK, bold=True, font=FONT_HEAD)

    rows = d["events"][:6]
    ty = 1.5
    row_h = 0.9
    for r in rows:
        _card(s, 0.6, ty, 11.3, row_h - 0.1, WHITE, line_color=RGBColor(0xE3, 0xE7, 0xF1))
        _box(s, 0.85, ty + 0.08, 8.3, 0.3, r["title"], 12.5, INK, bold=True)
        _box(s, 0.85, ty + 0.4, 8.3, 0.3, r["detail"], 10, MUTED)
        _box(s, 0.85, ty + 0.65, 8.3, 0.2, r["sources"], 9, SLATE, italic=True)
        color = GREEN if r["score"] >= 8 else (AMBER if r["score"] >= 5 else RED)
        circle = s.shapes.add_shape(9, Inches(10.75), Inches(ty + (row_h - 0.1) / 2 - 0.3), Inches(0.6), Inches(0.6))  # 9 = oval
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
        circle.shadow.inherit = False
        tf = circle.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(r["score"])
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = WHITE
        ty += row_h


def _slide_trend(s, d):
    _bg(s, PAPER)
    _box(s, 0.6, 0.4, 8, 0.35, "TREND ANALYSIS", 12, AMBER, bold=True)
    _box(s, 0.6, 0.72, 12.1, 0.55, "Threat-Level Trajectory", 24, INK, bold=True, font=FONT_HEAD)

    hist = d["threat_index_history"][-30:]
    chart_data = CategoryChartData()
    chart_data.categories = [h["date"][5:] for h in hist]  # MM-DD
    chart_data.add_series("Threat Index", [h["value"] for h in hist])
    x, y, cx, cy = Inches(0.6), Inches(1.5), Inches(8.1), Inches(5.2)
    gframe = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, x, y, cx, cy, chart_data)
    chart = gframe.chart
    chart.has_legend = False
    chart.has_title = False
    val_axis = chart.value_axis
    val_axis.minimum_scale = 0
    val_axis.maximum_scale = 10
    plot = chart.plots[0]
    plot.series[0].format.line.color.rgb = AMBER
    plot.series[0].format.line.width = Pt(2.5)

    _card(s, 8.95, 1.5, 3.8, 5.2, NAVY)
    _box(s, 9.2, 1.7, 3.3, 0.3, "KEY SIGNALS THIS WEEK", 11, ICE, bold=True)
    ky = 2.1
    for sig in d["key_signals"][:5]:
        _box(s, 9.2, ky, 3.35, 0.7, sig, 10.5, WHITE)
        ky += 0.75


def _slide_narrative_stability(s, d):
    _bg(s, PAPER)
    _box(s, 0.6, 0.4, 8, 0.35, "NARRATIVE STABILITY", 12, AMBER, bold=True)
    _box(s, 0.6, 0.72, 12.1, 0.55, "24h Correction & Drift Tracker", 24, INK, bold=True, font=FONT_HEAD)
    _box(s, 0.6, 1.28, 12.1, 0.3, "Based on hourly checks against each watchlist thread's prior known status", 11, MUTED, italic=True)

    rows = d["narrative_stability"][:6]
    ty = 1.75
    row_h = 0.87
    for r in rows:
        _card(s, 0.6, ty, 11.3, row_h - 0.1, WHITE, line_color=RGBColor(0xE3, 0xE7, 0xF1))
        _box(s, 0.85, ty + 0.08, 7.5, 0.3, r["thread"], 12, INK, bold=True)
        _box(s, 0.85, ty + 0.4, 7.5, 0.35, r["summary"], 10, MUTED)
        badge_color = {"stable": GREEN, "corrected": AMBER, "shifted": RED, "new": AMBER}.get(r["status"], MUTED)
        _box(s, 9.3, ty + 0.18, 2.5, 0.4, r["status"].upper(), 12, badge_color, bold=True, align=PP_ALIGN.RIGHT)
        ty += row_h


def _slide_source_library(s, d):
    _bg(s, PAPER)
    _box(s, 0.6, 0.4, 8, 0.35, "SOURCE LIBRARY", 12, AMBER, bold=True)
    _box(s, 0.6, 0.72, 12.1, 0.55, "Daily-Weighted Source & Author Reference", 24, INK, bold=True, font=FONT_HEAD)

    tiers = d["tiers"]
    ty = 1.4
    for t in tiers:
        h = 0.72
        _card(s, 0.6, ty, 11.3, h - 0.1, WHITE, line_color=RGBColor(0xE3, 0xE7, 0xF1))
        _box(s, 0.9, ty, 3.3, h - 0.1, t["name"], 11.5, INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        _box(s, 4.3, ty, 1.0, h - 0.1, t["range"], 15, AMBER if "Excluded" not in t["name"] else RED, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        _box(s, 5.4, ty, 6.4, h - 0.1, t["examples"], 10, MUTED, anchor=MSO_ANCHOR.MIDDLE)
        ty += h

    _box(s, 0.6, ty + 0.15, 11.3, 0.3, "TRACKED AUTHORS - MOST CORRECTIONS-TRANSPARENT (30d)", 10.5, AMBER, bold=True)
    _box(s, 0.6, ty + 0.45, 11.3, 0.7, d.get("top_authors_summary", "Building history - check back after a few days of monitoring."), 10.5, MUTED)


def _slide_deception_watch(s, d):
    _bg(s, NAVY_DK)
    _box(s, 0.6, 0.45, 8, 0.35, "DECEPTION WATCH", 12, AMBER, bold=True)
    _box(s, 0.6, 0.78, 12.1, 0.6, "Kremlin Narrative & Disinformation Tracking", 24, WHITE, bold=True, font=FONT_HEAD)

    items = d["deception_items"][:4]
    card_w, card_h, gap_x, gap_y = 5.85, 2.35, 0.25, 0.25
    for i, item in enumerate(items):
        col, row = i % 2, i // 2
        x = 0.6 + col * (card_w + gap_x)
        y = 1.85 + row * (card_h + gap_y)
        _card(s, x, y, card_w, card_h, NAVY)
        _box(s, x + 0.3, y + 0.25, card_w - 0.6, 0.5, item["heading"], 13.5, WHITE, bold=True)
        _box(s, x + 0.3, y + 0.85, card_w - 0.6, card_h - 1.05, item["body"], 11, ICE)


def _slide_proposal(s, d):
    _bg(s, NAVY_DK)
    _box(s, 0.6, 0.45, 8, 0.35, "METHODOLOGY", 12, AMBER, bold=True)
    _box(s, 0.6, 0.78, 12.1, 0.6, "Proposed Methodology Change - Today", 24, WHITE, bold=True, font=FONT_HEAD)

    p = d["proposal"]
    _card(s, 0.6, 1.85, 12.1, 3.2, NAVY)
    _box(s, 0.9, 2.1, 11.5, 0.4, p["title"], 16, AMBER, bold=True)
    _box(s, 0.9, 2.6, 11.5, 1.3, p["rationale"], 12, WHITE)
    _box(s, 0.9, 4.0, 11.5, 0.9, "Change: " + p["change_detail"], 12, ICE, italic=True)

    _box(s, 0.6, 5.3, 12.1, 0.4, "To respond: reply to today's report with \"accept\" or \"deny\".", 13, WHITE, bold=True)
    _box(s, 0.6, 5.75, 12.1, 0.4, "Or run:  python src/daily_report.py --resolve accept   /   --resolve deny", 11.5, SLATE, italic=True)
    if d.get("prior_proposal_still_pending"):
        _box(s, 0.6, 6.3, 12.1, 0.4, "Note: yesterday's proposal is still unresolved - resolve it before a new one is generated.", 11, AMBER, italic=True)
